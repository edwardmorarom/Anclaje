from pathlib import Path

import pytest

from anclaje.fragmentos import fragmentar
from anclaje.ingesta import leer_archivo, leer_corpus
from anclaje.modelos import Pagina


def test_fragmentos_conservan_pagina_y_origen():
    pages = [Pagina("publicos/a.txt", 1, "publicos", "A" * 22), Pagina("publicos/a.txt", 2, "publicos", "B" * 16)]
    chunks = fragmentar(pages, 10, 3)
    assert {f.pagina for f in chunks} == {1, 2}
    assert all(len(f.texto) <= 10 and f.origen == "publicos" for f in chunks)
    assert all(set(f.texto) == ({"A"} if f.pagina == 1 else {"B"}) for f in chunks)
    assert chunks == fragmentar(pages, 10, 3)
    assert len({f.id for f in chunks}) == len(chunks)


@pytest.mark.parametrize("size,overlap", [(10, 10), (10, -1), (0, 0)])
def test_fragmentacion_invalida(size, overlap):
    with pytest.raises(ValueError):
        fragmentar([], size, overlap)


def test_pdf_generado_y_pagina_sin_texto(tmp_path):
    import pymupdf

    path = tmp_path / "generado.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((40, 50), "Media: suma dividida entre n.")
        pdf.new_page()
        pdf.new_page().insert_text((40, 50), "Mediana: valor central.")
        pdf.save(path)
    with pytest.warns(UserWarning, match="p. 2.*sin texto"):
        pages = leer_archivo(path, "publicos/generado.pdf", "publicos")
    assert len(pages) == 3 and pages[2].pagina == 3
    assert "Mediana" in pages[2].texto
    assert pages[1].texto == ""


@pytest.mark.parametrize("suffix", ["txt", "md"])
def test_texto_utf8_y_saltos_formulario(tmp_path, suffix):
    path = tmp_path / f"generado.{suffix}"
    path.write_text("Estadística\fPoblación", encoding="utf-8-sig")
    pages = leer_archivo(path, f"publicos/{path.name}", "publicos")
    assert [p.pagina for p in pages] == [1, 2]
    assert pages[0].texto == "Estadística"


def test_docx_generado_con_salto_y_tabla(tmp_path):
    from docx import Document

    path = tmp_path / "generado.docx"
    doc = Document()
    doc.add_paragraph("Primera página")
    doc.add_page_break()
    doc.add_paragraph("Segunda página")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Resultado 42"
    doc.save(path)
    with pytest.warns(UserWarning, match="DOCX estimadas"):
        pages = leer_archivo(path, "contraparte/generado.docx", "contraparte")
    assert len(pages) == 2
    assert "Primera" in pages[0].texto and "Segunda" not in pages[0].texto
    assert "Segunda" in pages[1].texto and "Resultado 42" in pages[1].texto


def test_procedencia_plantilla_y_aviso_por_archivo(tmp_path):
    folder = tmp_path / "publicos"
    folder.mkdir()
    (folder / "a.txt").write_text("Documento sintético.", encoding="utf-8")
    with pytest.warns(UserWarning) as records:
        pages = leer_corpus(tmp_path)
    assert len(pages) == 1
    assert any("no aparece" in str(r.message) for r in records)
    provenance = tmp_path / "PROCEDENCIA.md"
    assert "licencia o permiso" in provenance.read_text(encoding="utf-8")
    provenance.write_text("| publicos/a.txt | publicos | 2026-10-05 | Sintético | prueba |", encoding="utf-8")
    import warnings

    with warnings.catch_warnings(record=True) as records:
        leer_corpus(tmp_path)
    assert not records


def test_procedencia_no_acepta_coincidencia_parcial(tmp_path):
    (tmp_path / "publicos").mkdir()
    (tmp_path / "publicos" / "a.txt").write_text("Dato", encoding="utf-8")
    (tmp_path / "PROCEDENCIA.md").write_text("| publicos/otra-a.txt | publicos |", encoding="utf-8")
    with pytest.warns(UserWarning, match="no aparece"):
        leer_corpus(tmp_path)
