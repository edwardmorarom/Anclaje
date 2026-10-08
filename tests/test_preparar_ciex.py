import json

import pytest
from docx import Document

from scripts.preparar_ciex import preparar


def materiales(tmp_path):
    base = tmp_path / 'originales'
    for carpeta in ['edward/Fuentes_seleccionadas', 'Harold', 'natalia']:
        folder = base / carpeta
        folder.mkdir(parents=True)
        (folder / 'articulo.pdf').write_bytes(b'contenido de prueba')
    (base / 'edward/Fuentes_seleccionadas/README_fuentes.md').write_text('DOI declarado', encoding='utf-8')
    Document().save(base / 'edward/Estado_del_arte_Mercado_Publico_y_Mascoplan_solo_PDF.docx')
    banco = tmp_path / 'banco.docx'
    doc = Document()
    tabla = doc.add_table(rows=1, cols=4)
    for celda, texto in zip(tabla.rows[0].cells, ['#', 'Fuente', 'Pregunta', 'Respuesta']):
        celda.text = texto
    for celda, texto in zip(tabla.add_row().cells, ['1', 'Vacío de Información', '¿Cuánto?', 'No está en las fuentes.']):
        celda.text = texto
    doc.save(banco)
    return base, banco


def test_separa_evidencias_y_conserva_revision_manual(tmp_path):
    base, banco = materiales(tmp_path)
    destino = tmp_path / 'extraccion'
    preparar(base, banco, destino)
    pdfs = list((destino / 'corpus').rglob('*.pdf'))
    assert len(pdfs) == 1 and pdfs[0].parent.name == 'Edward'
    assert (destino / 'evidencias/A/Harold/articulo.pdf').exists()
    assert (destino / 'evidencias/B/Natalia/articulo.pdf').exists()
    assert not (destino / 'evaluacion/banco.csv').exists()
    assert json.loads((destino / 'evaluacion/mesa_inicial.json').read_text(encoding='utf-8'))[0]['Pregunta'] == '¿Cuánto?'
    borrador = destino / 'evaluacion/banco_pendiente.csv'
    borrador.write_text('revisión humana', encoding='utf-8')
    preparar(base, banco, destino)
    assert borrador.read_text(encoding='utf-8') == 'revisión humana'


def test_no_sobrescribe_un_original_distinto(tmp_path):
    base, banco = materiales(tmp_path)
    destino = tmp_path / 'extraccion'
    preparar(base, banco, destino)
    (base / 'edward/Fuentes_seleccionadas/articulo.pdf').write_bytes(b'version nueva')
    with pytest.raises(ValueError, match='Destino distinto'):
        preparar(base, banco, destino)
    assert (destino / 'corpus/publicos/Edward/articulo.pdf').read_bytes() == b'contenido de prueba'
